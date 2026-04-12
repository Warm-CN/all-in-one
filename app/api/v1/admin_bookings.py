from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import date, timedelta
import io
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.database import get_db
from app.core.dependencies import require_admin
from app.models.room_booking import RoomBooking
from app.models.user import User
from app.schemas.response import success_response

router = APIRouter()

@router.get("/range", summary="获取时间范围内的预约")
async def get_bookings_range(
    start_date: date,
    end_date: date,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    bookings = db.query(RoomBooking).filter(
        RoomBooking.booking_date >= start_date,
        RoomBooking.booking_date <= end_date
    ).order_by(RoomBooking.booking_date.desc(), RoomBooking.start_time).all()
    
    data = []
    for booking in bookings:
        data.append({
            "id": booking.id,
            "user_id": booking.user_id,
            "booking_date": booking.booking_date.isoformat(),
            "start_time": booking.start_time.strftime("%H:%M"),
            "end_time": booking.end_time.strftime("%H:%M"),
            "num_people": booking.num_people,
            "remarks": booking.remarks,
            "user_name": booking.user.full_name if booking.user else "Unknown",
            "user_dept": booking.user.department if booking.user else "Unknown",
            "created_at": booking.created_at.isoformat()
        })
        
    return success_response(data=data)

@router.get("/export", summary="导出月度预约记录")
async def export_bookings(
    year: int,
    month: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    # 计算当月的起止日期
    start_date = date(year, month, 1)
    if month == 12:
        end_date = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        end_date = date(year, month + 1, 1) - timedelta(days=1)
        
    bookings = db.query(RoomBooking).filter(
        RoomBooking.booking_date >= start_date,
        RoomBooking.booking_date <= end_date
    ).order_by(RoomBooking.booking_date, RoomBooking.start_time).all()
    
    # 创建 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}年{month}月"
    
    # 表头
    headers = ["预约日期", "开始时间", "结束时间", "预约人", "部门", "参与人数", "备注", "提交时间"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        
    # 数据
    for row, booking in enumerate(bookings, 2):
        ws.cell(row=row, column=1, value=booking.booking_date.strftime("%Y-%m-%d"))
        ws.cell(row=row, column=2, value=booking.start_time.strftime("%H:%M"))
        ws.cell(row=row, column=3, value=booking.end_time.strftime("%H:%M"))
        ws.cell(row=row, column=4, value=booking.user.full_name if booking.user else "Unknown")
        ws.cell(row=row, column=5, value=booking.user.department if booking.user else "Unknown")
        ws.cell(row=row, column=6, value=booking.num_people)
        ws.cell(row=row, column=7, value=booking.remarks or "")
        ws.cell(row=row, column=8, value=booking.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        
    # 调整列宽
    widths = [15, 10, 10, 15, 15, 10, 30, 20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # 返回文件流 - 注意这里文件名如果包含中文可能需要URL编码，现代浏览器通常处理得好
    # 使用 ASCII 文件名以防万一
    filename = f"bookings_{year}_{month}.xlsx"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        output, 
        headers=headers, 
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
