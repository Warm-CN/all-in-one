"""
竞赛队伍模块路由
公开端 + 成员端 + 管理端
"""
from datetime import date, datetime
from io import BytesIO
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, field_validator
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import require_admin, require_member
from app.models.user import User
from app.models.team import Team, TeamStatus
from app.models.team_member import TeamMember
from app.models.topic_option import TopicOption
from app.models.team_channel_config import TeamChannelConfig
from app.models.inspection_assignment import InspectionAssignment
from app.schemas.response import success_response, error_response

router = APIRouter()

DEFAULT_MODULE_KEY = "wireless_cup_2026"


class TeamMemberPayload(BaseModel):
    name: str
    student_id: str
    phone: str
    email: Optional[str] = None
    college: str
    major_class: str

    @field_validator("name", "student_id", "phone", "college", "major_class")
    @classmethod
    def validate_required_text(cls, value: str) -> str:
        value = (value or "").strip()
        if not value:
            raise ValueError("队员信息不完整，请补齐姓名/学号/手机号")
        return value

    @field_validator("student_id")
    @classmethod
    def validate_student_id_length(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) != 12:
            raise ValueError("学号默认要求 12 位")
        return value


class TeamCreateRequest(BaseModel):
    team_name: str
    competition_track: Optional[str] = None
    captain_name: str
    captain_student_id: str
    captain_phone: str
    captain_email: Optional[str] = None
    captain_college: str
    captain_major_class: str
    topic_id: Optional[int] = None
    module_key: str = DEFAULT_MODULE_KEY
    members: List[TeamMemberPayload] = []

    @field_validator(
        "team_name",
        "captain_name",
        "captain_student_id",
        "captain_phone",
        "captain_college",
        "captain_major_class",
        "module_key",
    )
    @classmethod
    def validate_main_required_text(cls, value: str) -> str:
        value = (value or "").strip()
        if not value:
            raise ValueError("队长与队伍关键信息不能为空")
        return value

    @field_validator("captain_student_id")
    @classmethod
    def validate_captain_student_id_length(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) != 12:
            raise ValueError("学号默认要求 12 位")
        return value


class TeamUpdateRequest(BaseModel):
    captain_student_id: str
    captain_phone: str
    team_name: str
    competition_track: Optional[str] = None
    captain_name: str
    captain_email: Optional[str] = None
    captain_college: str
    captain_major_class: str
    members: List[TeamMemberPayload] = []

    @field_validator("team_name", "captain_name", "captain_phone", "captain_college", "captain_major_class")
    @classmethod
    def validate_required_text(cls, value: str) -> str:
        value = (value or "").strip()
        if not value:
            raise ValueError("队长与队伍关键信息不能为空")
        return value

    @field_validator("captain_student_id")
    @classmethod
    def validate_captain_student_id_length(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) != 12:
            raise ValueError("学号默认要求 12 位")
        return value


class TeamTopicUpdateRequest(BaseModel):
    captain_student_id: str
    topic_id: int
    module_key: str = DEFAULT_MODULE_KEY

    @field_validator("captain_student_id")
    @classmethod
    def validate_captain_student_id_length(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) != 12:
            raise ValueError("学号默认要求 12 位")
        return value


class TeamDeleteRequest(BaseModel):
    captain_student_id: str
    captain_phone: str
    delete_reason: str
    confirm_self_operation: bool
    confirm_members_informed: bool
    module_key: str = DEFAULT_MODULE_KEY

    @field_validator("captain_student_id")
    @classmethod
    def validate_captain_student_id_length(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) != 12:
            raise ValueError("学号默认要求 12 位")
        return value

    @field_validator("delete_reason")
    @classmethod
    def validate_delete_reason(cls, value: str) -> str:
        value = (value or "").strip()
        if len(value) < 5:
            raise ValueError("删除理由至少 5 个字符")
        return value


class TeamBatchDeleteRequest(BaseModel):
    team_ids: List[int]


class TeamBatchTopicRequest(BaseModel):
    team_ids: List[int]
    topic_id: int


class TeamInspectionUpdateRequest(BaseModel):
    first_inspection_time: Optional[datetime] = None
    first_inspection_location: Optional[str] = None
    first_inspector: Optional[str] = None
    first_notes: Optional[str] = None
    second_inspection_time: Optional[datetime] = None
    second_inspection_location: Optional[str] = None
    second_inspector: Optional[str] = None
    second_notes: Optional[str] = None

    # 兼容旧字段
    inspection_time: Optional[datetime] = None
    inspection_location: Optional[str] = None
    inspector: Optional[str] = None
    notes: Optional[str] = None


class TeamBatchInspectionUpdateRequest(BaseModel):
    team_ids: List[int]
    first_inspection_time: Optional[datetime] = None
    first_inspection_location: Optional[str] = None
    first_inspector: Optional[str] = None
    first_notes: Optional[str] = None
    second_inspection_time: Optional[datetime] = None
    second_inspection_location: Optional[str] = None
    second_inspector: Optional[str] = None
    second_notes: Optional[str] = None

    # 兼容旧字段
    inspection_time: Optional[datetime] = None
    inspection_location: Optional[str] = None
    inspector: Optional[str] = None
    notes: Optional[str] = None

    check_duplicate: bool = True


class TeamChannelConfigUpdateRequest(BaseModel):
    module_key: str = DEFAULT_MODULE_KEY
    signup_open: Optional[bool] = None
    topic_open: Optional[bool] = None
    info_update_open: Optional[bool] = None
    signup_close_at: Optional[datetime] = None
    topic_close_at: Optional[datetime] = None
    info_update_close_at: Optional[datetime] = None


class TopicUpsertRequest(BaseModel):
    code: str
    title: str
    description: Optional[str] = None
    is_active: bool = True
    max_teams: Optional[int] = None


def _normalize_excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    return str(value).strip()


def _parse_excel_datetime(value) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    text = _normalize_excel_text(value)
    if not text:
        return None

    text = text.replace("T", " ").replace("Z", "").strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _header_index_map(sheet) -> dict:
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
    return {
        _normalize_excel_text(header): index
        for index, header in enumerate(headers)
        if _normalize_excel_text(header)
    }


def _first_cell_value(row, header_map: dict, names: List[str]):
    for name in names:
        index = header_map.get(name)
        if index is not None and index < len(row):
            return row[index]
    return None


def _find_topic_by_excel_value(db: Session, module_key: str, value) -> Optional[TopicOption]:
    text = _normalize_excel_text(value)
    if not text:
        return None

    query = db.query(TopicOption)
    if module_key:
        query = query.filter(TopicOption.module_key == module_key)

    if text.isdigit():
        topic = query.filter(TopicOption.id == int(text)).first()
        if topic:
            return topic

    candidates = [text]
    if " - " in text:
        left, right = text.split(" - ", 1)
        candidates.extend([left.strip(), right.strip()])
    if "-" in text:
        left, right = text.split("-", 1)
        candidates.extend([left.strip(), right.strip()])

    for candidate in [item for item in candidates if item]:
        topic = query.filter(TopicOption.code == candidate).first()
        if topic:
            return topic
        topic = query.filter(TopicOption.title == candidate).first()
        if topic:
            return topic

    return None


def _ensure_channel_config(db: Session, module_key: str) -> TeamChannelConfig:
    cfg = db.query(TeamChannelConfig).filter(TeamChannelConfig.module_key == module_key).first()
    if cfg:
        return cfg

    cfg = TeamChannelConfig(
        module_key=module_key,
        signup_open=False,
        topic_open=False,
        info_update_open=True,
        signup_close_at=None,
        topic_close_at=None,
        info_update_close_at=None,
    )
    db.add(cfg)
    db.commit()
    db.refresh(cfg)
    return cfg


def _team_to_dict(team: Team) -> dict:
    topic_title = team.topic.title if team.topic else None
    assignment = team.inspection_assignment
    return {
        "id": team.id,
        "team_name": team.team_name,
        "competition_track": team.competition_track,
        "module_key": team.module_key,
        "captain_name": team.captain_name,
        "captain_student_id": team.captain_student_id,
        "captain_phone": team.captain_phone,
        "captain_email": team.captain_email,
        "captain_college": team.captain_college,
        "captain_major_class": team.captain_major_class,
        "topic_id": team.topic_id,
        "topic_title": topic_title,
        "status": team.status.value if isinstance(team.status, TeamStatus) else str(team.status),
        "notes": team.notes,
        "delete_requested": bool(team.delete_requested),
        "delete_requested_at": team.delete_requested_at.isoformat() if team.delete_requested_at else None,
        "delete_reason": team.delete_reason,
        "members": [
            {
                "id": m.id,
                "name": m.name,
                "student_id": m.student_id,
                "phone": m.phone,
                "email": m.email,
                "college": m.college,
                "major_class": m.major_class,
            }
            for m in team.members
        ],
        "inspection": {
            "first_inspection_time": assignment.first_inspection_time.isoformat() if assignment and assignment.first_inspection_time else None,
            "first_inspection_location": assignment.first_inspection_location if assignment else None,
            "first_inspector": assignment.first_inspector if assignment else None,
            "first_notes": assignment.first_notes if assignment else None,
            "second_inspection_time": assignment.second_inspection_time.isoformat() if assignment and assignment.second_inspection_time else None,
            "second_inspection_location": assignment.second_inspection_location if assignment else None,
            "second_inspector": assignment.second_inspector if assignment else None,
            "second_notes": assignment.second_notes if assignment else None,
            # 兼容旧字段
            "inspection_time": assignment.inspection_time.isoformat() if assignment and assignment.inspection_time else None,
            "inspection_location": assignment.inspection_location if assignment else None,
            "inspector": assignment.inspector if assignment else None,
            "notes": assignment.notes if assignment else None,
        },
        "created_at": team.created_at.isoformat(),
        "updated_at": team.updated_at.isoformat(),
    }


def _validate_member_uniqueness(payload_members: List[TeamMemberPayload]) -> Optional[str]:
    if len(payload_members) > 2:
        return "队员最多 2 人（不含队长）"

    sids = [m.student_id for m in payload_members]
    if len(set(sids)) != len(sids):
        return "队员学号重复，请检查后重试"
    return None


def _assert_topic_available(db: Session, topic_id: int, module_key: Optional[str] = None) -> Optional[str]:
    query = db.query(TopicOption).filter(TopicOption.id == topic_id, TopicOption.is_active == True)
    if module_key:
        query = query.filter(TopicOption.module_key == module_key)

    topic = query.first()
    if not topic:
        return "题目不存在或不可选"

    if topic.max_teams is not None:
        used_count = db.query(Team).filter(Team.topic_id == topic_id).count()
        if used_count >= topic.max_teams:
            return "该题目已达到队伍上限"
    return None


@router.get("/topics", response_model=dict, summary="获取题目列表")
async def list_topics(only_active: bool = True, module_key: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(TopicOption)
    if only_active:
        query = query.filter(TopicOption.is_active == True)
    if module_key:
        query = query.filter(TopicOption.module_key == module_key)

    topics = query.order_by(TopicOption.id.asc()).all()
    return success_response(
        data=[
            {
                "id": t.id,
                "code": t.code,
                "title": t.title,
                "description": t.description,
                "is_active": t.is_active,
                "max_teams": t.max_teams,
            }
            for t in topics
        ],
        msg=f"获取成功，共 {len(topics)} 条",
    )


@router.post("/topics", response_model=dict, summary="管理员新增题目")
async def create_topic(
    req: TopicUpsertRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    existing_code = db.query(TopicOption).filter(TopicOption.code == req.code).first()
    if existing_code:
        return error_response(400, "题目编号已存在")
    existing_title = db.query(TopicOption).filter(TopicOption.title == req.title).first()
    if existing_title:
        return error_response(400, "题目名称已存在")

    topic = TopicOption(
        code=req.code.strip(),
        title=req.title.strip(),
        description=req.description,
        is_active=req.is_active,
        max_teams=req.max_teams,
    )
    db.add(topic)
    db.commit()
    db.refresh(topic)
    return success_response(data={"topic_id": topic.id}, msg="新增成功")


@router.put("/topics/{topic_id}", response_model=dict, summary="管理员修改题目")
async def update_topic(
    topic_id: int,
    req: TopicUpsertRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    topic = db.query(TopicOption).filter(TopicOption.id == topic_id).first()
    if not topic:
        return error_response(404, "题目不存在")

    duplicate_code = db.query(TopicOption).filter(TopicOption.code == req.code, TopicOption.id != topic_id).first()
    if duplicate_code:
        return error_response(400, "题目编号已存在")
    duplicate_title = db.query(TopicOption).filter(TopicOption.title == req.title, TopicOption.id != topic_id).first()
    if duplicate_title:
        return error_response(400, "题目名称已存在")

    topic.code = req.code.strip()
    topic.title = req.title.strip()
    topic.description = req.description
    topic.is_active = req.is_active
    topic.max_teams = req.max_teams
    db.commit()
    return success_response(msg="修改成功")


@router.get("/signup_status", response_model=dict, summary="报名通道状态")
async def signup_status(module_key: str = DEFAULT_MODULE_KEY, db: Session = Depends(get_db)):
    cfg = _ensure_channel_config(db, module_key)
    return success_response(
        data={
            "module_key": cfg.module_key,
            "signup_open": cfg.signup_open,
            "signup_close_at": cfg.signup_close_at.isoformat() if cfg.signup_close_at else None,
        },
        msg="获取成功",
    )


@router.get("/topic_status", response_model=dict, summary="选题通道状态")
async def topic_status(module_key: str = DEFAULT_MODULE_KEY, db: Session = Depends(get_db)):
    cfg = _ensure_channel_config(db, module_key)
    return success_response(
        data={
            "module_key": cfg.module_key,
            "topic_open": cfg.topic_open,
            "topic_close_at": cfg.topic_close_at.isoformat() if cfg.topic_close_at else None,
        },
        msg="获取成功",
    )


@router.get("/team_update_status", response_model=dict, summary="队伍信息修改通道状态")
async def team_update_status(module_key: str = DEFAULT_MODULE_KEY, db: Session = Depends(get_db)):
    cfg = _ensure_channel_config(db, module_key)
    return success_response(
        data={
            "module_key": cfg.module_key,
            "info_update_open": cfg.info_update_open,
            "info_update_close_at": cfg.info_update_close_at.isoformat() if cfg.info_update_close_at else None,
        },
        msg="获取成功",
    )


@router.post("/teams", response_model=dict, summary="公开端提交队伍报名")
async def create_team(req: TeamCreateRequest, db: Session = Depends(get_db)):
    cfg = _ensure_channel_config(db, req.module_key)
    now = datetime.now()
    if not cfg.signup_open:
        return error_response(400, "报名通道未开启")
    if cfg.signup_close_at and now > cfg.signup_close_at:
        return error_response(400, "报名已截止")

    if req.topic_id is not None:
        topic_error = _assert_topic_available(db, req.topic_id, req.module_key)
        if topic_error:
            return error_response(400, topic_error)

    duplicate_msg = _validate_member_uniqueness(req.members)
    if duplicate_msg:
        return error_response(400, duplicate_msg)

    if any(m.student_id == req.captain_student_id for m in req.members):
        return error_response(400, "队长学号不能出现在队员列表中")

    captain_exists = (
        db.query(Team)
        .filter(Team.captain_student_id == req.captain_student_id, Team.module_key == req.module_key)
        .first()
    )
    if captain_exists:
        return error_response(400, "该队长学号已报名")

    member_sids = [m.student_id for m in req.members]
    if member_sids:
        conflict_member = (
            db.query(TeamMember)
            .join(Team, TeamMember.team_id == Team.id)
            .filter(TeamMember.student_id.in_(member_sids), Team.module_key == req.module_key)
            .first()
        )
        if conflict_member:
            return error_response(400, f"学号 {conflict_member.student_id} 已在其他队伍")
        conflict_captain = (
            db.query(Team)
            .filter(Team.captain_student_id.in_(member_sids), Team.module_key == req.module_key)
            .first()
        )
        if conflict_captain:
            return error_response(400, f"学号 {conflict_captain.captain_student_id} 已作为队长报名")

    team = Team(
        team_name=req.team_name,
        competition_track=req.competition_track,
        module_key=req.module_key,
        captain_name=req.captain_name,
        captain_student_id=req.captain_student_id,
        captain_phone=req.captain_phone,
        captain_email=req.captain_email,
        captain_college=req.captain_college,
        captain_major_class=req.captain_major_class,
        topic_id=req.topic_id,
        status=TeamStatus.SUBMITTED,
    )
    db.add(team)
    db.flush()

    for member in req.members:
        db.add(
            TeamMember(
                team_id=team.id,
                name=member.name,
                student_id=member.student_id,
                phone=member.phone,
                email=member.email,
                college=member.college,
                major_class=member.major_class,
            )
        )

    db.commit()
    db.refresh(team)
    return success_response(data={"team_id": team.id}, msg="报名成功")


@router.get("/teams/by_sid/{sid}", response_model=dict, summary="公开端按学号查询队伍")
async def get_team_by_sid(sid: str, module_key: Optional[str] = None, db: Session = Depends(get_db)):
    team_query = db.query(Team).filter(Team.captain_student_id == sid)
    if module_key:
        team_query = team_query.filter(Team.module_key == module_key)
    team = team_query.first()

    if not team:
        member_query = db.query(TeamMember).filter(TeamMember.student_id == sid)
        if module_key:
            member_query = member_query.join(Team, TeamMember.team_id == Team.id).filter(Team.module_key == module_key)
        member = member_query.first()
        if member:
            team = db.query(Team).filter(Team.id == member.team_id).first()

    if not team:
        return error_response(404, "未查询到对应队伍")

    return success_response(data=_team_to_dict(team), msg="查询成功")


@router.put("/teams/{team_id:int}", response_model=dict, summary="公开端修改队伍信息")
async def update_team(team_id: int, req: TeamUpdateRequest, db: Session = Depends(get_db)):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")
    if team.delete_requested:
        return error_response(400, "队伍已提交删除申请，等待管理员确认后不可修改")

    cfg = _ensure_channel_config(db, team.module_key or DEFAULT_MODULE_KEY)
    now = datetime.now()
    if not cfg.info_update_open:
        return error_response(400, "队伍信息修改通道未开启")
    if cfg.info_update_close_at and now > cfg.info_update_close_at:
        return error_response(400, "队伍信息修改已截止")

    if team.captain_student_id != req.captain_student_id or team.captain_phone != req.captain_phone:
        return error_response(403, "仅允许队长修改本队信息")

    duplicate_msg = _validate_member_uniqueness(req.members)
    if duplicate_msg:
        return error_response(400, duplicate_msg)
    if any(m.student_id == req.captain_student_id for m in req.members):
        return error_response(400, "队长学号不能出现在队员列表中")

    member_sids = [m.student_id for m in req.members]
    if member_sids:
        conflict_member = (
            db.query(TeamMember)
            .join(Team, TeamMember.team_id == Team.id)
            .filter(
                TeamMember.student_id.in_(member_sids),
                TeamMember.team_id != team_id,
                Team.module_key == (team.module_key or DEFAULT_MODULE_KEY),
            )
            .first()
        )
        if conflict_member:
            return error_response(400, f"学号 {conflict_member.student_id} 已在其他队伍")

        conflict_captain = (
            db.query(Team)
            .filter(
                Team.captain_student_id.in_(member_sids),
                Team.id != team_id,
                Team.module_key == (team.module_key or DEFAULT_MODULE_KEY),
            )
            .first()
        )
        if conflict_captain:
            return error_response(400, f"学号 {conflict_captain.captain_student_id} 已作为队长报名")

    team.team_name = req.team_name
    team.competition_track = req.competition_track
    team.captain_name = req.captain_name
    team.captain_email = req.captain_email
    team.captain_college = req.captain_college
    team.captain_major_class = req.captain_major_class

    db.query(TeamMember).filter(TeamMember.team_id == team_id).delete()
    for member in req.members:
        db.add(
            TeamMember(
                team_id=team.id,
                name=member.name,
                student_id=member.student_id,
                phone=member.phone,
                email=member.email,
                college=member.college,
                major_class=member.major_class,
            )
        )

    db.commit()
    return success_response(msg="队伍信息修改成功")


@router.delete("/teams/{team_id:int}", response_model=dict, summary="公开端删除队伍")
async def delete_team(team_id: int, req: TeamDeleteRequest, db: Session = Depends(get_db)):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")

    cfg = _ensure_channel_config(db, req.module_key)
    now = datetime.now()
    if not cfg.info_update_open:
        return error_response(400, "队伍信息修改通道未开启，暂不可删除")
    if cfg.info_update_close_at and now > cfg.info_update_close_at:
        return error_response(400, "队伍信息修改已截止，暂不可删除")

    if team.captain_student_id != req.captain_student_id or team.captain_phone != req.captain_phone:
        return error_response(403, "仅允许队长删除本队")

    if not req.confirm_self_operation:
        return error_response(400, "请确认由本人操作后再提交")
    if not req.confirm_members_informed:
        return error_response(400, "请确认队员已知情后再提交")

    delete_reason = (req.delete_reason or "").strip()
    if len(delete_reason) < 5:
        return error_response(400, "删除理由至少 5 个字符")

    if team.delete_requested:
        return error_response(400, "该队伍已提交删除申请，请勿重复提交")

    team.delete_requested = True
    team.delete_requested_at = datetime.now()
    team.delete_reason = delete_reason
    db.commit()
    return success_response(msg="删除申请已提交，待管理员确认后才会彻底删除")


@router.put("/teams/{team_id:int}/topic", response_model=dict, summary="公开端修改队伍选题")
async def update_team_topic(team_id: int, req: TeamTopicUpdateRequest, db: Session = Depends(get_db)):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")
    if team.delete_requested:
        return error_response(400, "队伍已提交删除申请，等待管理员确认后不可修改选题")

    if team.captain_student_id != req.captain_student_id:
        return error_response(403, "仅允许队长修改本队选题")

    cfg = _ensure_channel_config(db, req.module_key)
    now = datetime.now()
    if not cfg.topic_open:
        return error_response(400, "选题通道未开启")
    if cfg.topic_close_at and now > cfg.topic_close_at:
        return error_response(400, "选题已截止")

    topic_error = _assert_topic_available(db, req.topic_id, req.module_key)
    if topic_error:
        return error_response(400, topic_error)

    team.topic_id = req.topic_id
    db.commit()
    return success_response(msg="选题修改成功")


@router.get("/teams", response_model=dict, summary="成员/管理员查看队伍列表")
async def list_teams(
    keyword: Optional[str] = None,
    topic_id: Optional[int] = None,
    module_key: Optional[str] = None,
    delete_requested: Optional[bool] = None,
    current_user: User = Depends(require_member),
    db: Session = Depends(get_db),
):
    query = db.query(Team)
    if keyword:
        query = query.filter(
            Team.team_name.contains(keyword)
            | Team.captain_name.contains(keyword)
            | Team.captain_student_id.contains(keyword)
        )
    if topic_id:
        query = query.filter(Team.topic_id == topic_id)
    if module_key:
        query = query.filter(Team.module_key == module_key)
    if delete_requested is not None:
        query = query.filter(Team.delete_requested == delete_requested)

    teams = query.order_by(Team.created_at.desc()).all()
    return success_response(data=[_team_to_dict(team) for team in teams], msg=f"获取成功，共 {len(teams)} 条")


@router.post("/teams/{team_id:int}/confirm-delete", response_model=dict, summary="管理员确认彻底删除队伍")
async def confirm_delete_team(
    team_id: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")
    if not team.delete_requested:
        return error_response(400, "该队伍尚未提交删除申请")

    db.delete(team)
    db.commit()
    return success_response(msg="已确认并彻底删除该队伍")


@router.get("/teams/export", summary="成员/管理员导出队伍数据")
async def export_teams(
    keyword: Optional[str] = None,
    topic_id: Optional[int] = None,
    module_key: Optional[str] = None,
    delete_requested: Optional[bool] = None,
    current_user: User = Depends(require_member),
    db: Session = Depends(get_db),
):
    query = db.query(Team)
    if keyword:
        query = query.filter(
            Team.team_name.contains(keyword)
            | Team.captain_name.contains(keyword)
            | Team.captain_student_id.contains(keyword)
        )
    if topic_id:
        query = query.filter(Team.topic_id == topic_id)
    if module_key:
        query = query.filter(Team.module_key == module_key)
    if delete_requested is not None:
        query = query.filter(Team.delete_requested == delete_requested)

    teams = query.order_by(Team.created_at.desc()).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "队伍列表"
    headers = [
        "队伍ID",
        "队伍名称",
        "赛道",
        "队长姓名",
        "队长学号",
        "队长手机号",
        "队长邮箱",
        "队长学院",
        "队长专业班级",
        "队员1-姓名",
        "队员1-学号",
        "队员1-手机号",
        "队员1-邮箱",
        "队员1-学院",
        "队员1-专业班级",
        "队员2-姓名",
        "队员2-学号",
        "队员2-手机号",
        "队员2-邮箱",
        "队员2-学院",
        "队员2-专业班级",
        "选题",
        "状态",
        "队员人数",
        "一验时间",
        "一验地点",
        "一验备注",
        "二验时间",
        "二验地点",
        "二验备注",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for team in teams:
        assignment = team.inspection_assignment
        ordered_members = sorted(team.members, key=lambda m: m.id)
        member1 = ordered_members[0] if len(ordered_members) > 0 else None
        member2 = ordered_members[1] if len(ordered_members) > 1 else None
        sheet.append(
            [
                team.id,
                team.team_name,
                team.competition_track or "",
                team.captain_name,
                team.captain_student_id,
                team.captain_phone,
                team.captain_email or "",
                team.captain_college or "",
                team.captain_major_class or "",
                member1.name if member1 else "",
                member1.student_id if member1 else "",
                member1.phone if member1 else "",
                member1.email if member1 else "",
                member1.college if member1 else "",
                member1.major_class if member1 else "",
                member2.name if member2 else "",
                member2.student_id if member2 else "",
                member2.phone if member2 else "",
                member2.email if member2 else "",
                member2.college if member2 else "",
                member2.major_class if member2 else "",
                team.topic.title if team.topic else "",
                team.status.value if isinstance(team.status, TeamStatus) else str(team.status),
                len(team.members),
                assignment.first_inspection_time.strftime("%Y-%m-%d %H:%M") if assignment and assignment.first_inspection_time else "",
                assignment.first_inspection_location if assignment else "",
                assignment.first_notes if assignment else "",
                assignment.second_inspection_time.strftime("%Y-%m-%d %H:%M") if assignment and assignment.second_inspection_time else "",
                assignment.second_inspection_location if assignment else "",
                assignment.second_notes if assignment else "",
            ]
        )

    widths = {
        "A": 10,
        "B": 20,
        "C": 16,
        "D": 12,
        "E": 16,
        "F": 16,
        "G": 24,
        "H": 20,
        "I": 22,
        "J": 12,
        "K": 16,
        "L": 16,
        "M": 24,
        "N": 20,
        "O": 22,
        "P": 12,
        "Q": 16,
        "R": 16,
        "S": 24,
        "T": 20,
        "U": 22,
        "V": 12,
        "W": 10,
        "X": 20,
        "Y": 20,
        "Z": 28,
        "AA": 20,
        "AB": 20,
        "AC": 28,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    example_sheet = workbook.create_sheet("填写示例")
    example_sheet.append(headers)
    example_sheet.append(
        [
            "请保留真实队伍ID",
            "示例队伍",
            "软件组",
            "张三",
            "202600000001",
            "13800000000",
            "example@example.com",
            "信息学院",
            "计科2601",
            "李四",
            "202600000002",
            "13800000001",
            "",
            "信息学院",
            "计科2601",
            "",
            "",
            "",
            "",
            "",
            "",
            "示例题目",
            "submitted",
            1,
            "2026-06-20 19:00",
            "创新楼 A101",
            "请携带作品与答辩材料",
            "2026-06-27 19:00",
            "创新楼 A102",
            "二验备注示例",
        ]
    )
    for cell in example_sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in widths.items():
        example_sheet.column_dimensions[col].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = f"teams_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("/teams/import", response_model=dict, summary="管理员导入队伍报名信息")
async def import_teams(
    module_key: str = DEFAULT_MODULE_KEY,
    file: UploadFile = File(...),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return error_response(400, "请上传 xlsx 文件")

    content = await file.read()
    if not content:
        return error_response(400, "上传文件为空")

    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        return error_response(400, "Excel 文件解析失败，请检查文件格式")

    sheet = workbook.active
    header_map = _header_index_map(sheet)
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if not rows:
        return error_response(400, "Excel 中没有可导入的数据")

    member_headers_exist = any(
        name in header_map
        for name in ["队员1-姓名", "队员1姓名", "队员2-姓名", "队员2姓名"]
    )

    created = 0
    updated = 0
    unchanged = 0
    skipped = 0
    errors = []

    for row_number, row in enumerate(rows, start=2):
        if not any(_normalize_excel_text(cell) for cell in row):
            continue

        team_id_text = _normalize_excel_text(_first_cell_value(row, header_map, ["队伍ID", "队伍编号", "team_id", "ID"]))
        team_name = _normalize_excel_text(_first_cell_value(row, header_map, ["队伍名称", "team_name"]))
        competition_track = _normalize_excel_text(_first_cell_value(row, header_map, ["赛道", "组别", "competition_track"]))
        captain_name = _normalize_excel_text(_first_cell_value(row, header_map, ["队长姓名", "captain_name"]))
        captain_student_id = _normalize_excel_text(_first_cell_value(row, header_map, ["队长学号", "captain_student_id"]))
        captain_phone = _normalize_excel_text(_first_cell_value(row, header_map, ["队长手机号", "队长电话", "captain_phone"]))
        captain_email = _normalize_excel_text(_first_cell_value(row, header_map, ["队长邮箱", "captain_email"]))
        captain_college = _normalize_excel_text(_first_cell_value(row, header_map, ["队长学院", "captain_college"]))
        captain_major_class = _normalize_excel_text(_first_cell_value(row, header_map, ["队长专业班级", "队长班级", "captain_major_class"]))

        if not team_name or not captain_name or not captain_student_id or not captain_phone:
            skipped += 1
            errors.append(f"第 {row_number} 行：队伍名称、队长姓名、队长学号、队长手机号为必填")
            continue

        team = None
        if team_id_text:
            try:
                team_id = int(float(team_id_text))
                team = db.query(Team).filter(Team.id == team_id, Team.module_key == module_key).first()
            except ValueError:
                errors.append(f"第 {row_number} 行：队伍ID格式不正确")

        if not team:
            team = db.query(Team).filter(
                Team.captain_student_id == captain_student_id,
                Team.module_key == module_key,
            ).first()

        topic = _find_topic_by_excel_value(
            db,
            module_key,
            _first_cell_value(row, header_map, ["选题", "题目", "topic_title", "topic_id"]),
        )

        is_new = team is None
        row_changed = False

        if is_new:
            team = Team(
                team_name=team_name,
                competition_track=competition_track or None,
                module_key=module_key,
                captain_name=captain_name,
                captain_student_id=captain_student_id,
                captain_phone=captain_phone,
                captain_email=captain_email or None,
                captain_college=captain_college or None,
                captain_major_class=captain_major_class or None,
                topic_id=topic.id if topic else None,
                status=TeamStatus.SUBMITTED,
            )
            db.add(team)
            db.flush()
            row_changed = True
        else:
            updates = {
                "team_name": team_name,
                "competition_track": competition_track or None,
                "captain_name": captain_name,
                "captain_student_id": captain_student_id,
                "captain_phone": captain_phone,
                "captain_email": captain_email or None,
                "captain_college": captain_college or None,
                "captain_major_class": captain_major_class or None,
            }
            if topic:
                updates["topic_id"] = topic.id

            for field_name, field_value in updates.items():
                if getattr(team, field_name) != field_value:
                    setattr(team, field_name, field_value)
                    row_changed = True

        if member_headers_exist:
            member_payloads = []
            for index in [1, 2]:
                name = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-姓名", f"队员{index}姓名"]))
                student_id = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-学号", f"队员{index}学号"]))
                phone = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-手机号", f"队员{index}手机号", f"队员{index}-电话"]))
                email = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-邮箱", f"队员{index}邮箱"]))
                college = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-学院", f"队员{index}学院"]))
                major_class = _normalize_excel_text(_first_cell_value(row, header_map, [f"队员{index}-专业班级", f"队员{index}专业班级", f"队员{index}-班级"]))

                if not any([name, student_id, phone, email, college, major_class]):
                    continue
                if not name or not student_id or not phone:
                    errors.append(f"第 {row_number} 行：队员{index}姓名、学号、手机号不完整，已跳过该队员")
                    continue

                member_payloads.append({
                    "name": name,
                    "student_id": student_id,
                    "phone": phone,
                    "email": email or None,
                    "college": college or None,
                    "major_class": major_class or None,
                })

            existing_members = [
                {
                    "name": member.name,
                    "student_id": member.student_id,
                    "phone": member.phone,
                    "email": member.email,
                    "college": member.college,
                    "major_class": member.major_class,
                }
                for member in sorted(team.members, key=lambda item: item.id)
            ]
            if existing_members != member_payloads:
                db.query(TeamMember).filter(TeamMember.team_id == team.id).delete()
                for member_data in member_payloads:
                    db.add(TeamMember(team_id=team.id, **member_data))
                row_changed = True

        if is_new:
            created += 1
        elif row_changed:
            updated += 1
        else:
            unchanged += 1

    db.commit()
    return success_response(
        data={
            "created": created,
            "updated": updated,
            "unchanged": unchanged,
            "skipped": skipped,
            "errors": errors[:20],
        },
        msg=f"导入完成，新增 {created} 条，更新 {updated} 条，未变化 {unchanged} 条，跳过 {skipped} 条",
    )


@router.get("/teams/inspection-template", summary="下载验收信息导入模板（仅管理员）")
async def download_inspection_template(
    current_user: User = Depends(require_admin),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "验收信息导入"
    headers = ["队伍ID", "队长学号", "队伍名称", "一验时间", "一验地点", "一验备注", "二验时间", "二验地点", "二验备注"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    column_widths = {"A": 10, "B": 15, "C": 20, "D": 18, "E": 18, "F": 25, "G": 18, "H": 18, "I": 25}
    for col, width in column_widths.items():
        sheet.column_dimensions[col].width = width

    help_sheet = workbook.create_sheet("填写说明")
    help_sheet.append(["字段", "说明"])
    help_sheet.append(["队伍ID/队长学号/队伍名称", "用于匹配队伍，填任一项即可"])
    help_sheet.append(["一验时间/地点/备注", "第一次验收安排，如：2026-07-15 14:00 / 教学楼A301 / 准备PPT"])
    help_sheet.append(["二验时间/地点/备注", "第二次验收安排"])
    help_sheet.append(["导入规则", "只更新 Excel 中填写了内容的单元格；留空不会清空已有安排"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    help_sheet.column_dimensions["A"].width = 28
    help_sheet.column_dimensions["B"].width = 64

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = "验收信息导入模板.xlsx"
    encoded_filename = quote(filename)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


@router.post("/teams/inspection-import", response_model=dict, summary="管理员导入队伍验收安排")
async def import_team_inspections(
    module_key: Optional[str] = None,
    file: UploadFile = File(...),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return error_response(400, "请上传 xlsx 文件")

    content = await file.read()
    if not content:
        return error_response(400, "上传文件为空")

    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        return error_response(400, "Excel 文件解析失败，请检查文件格式")

    sheet = workbook.active
    header_map = _header_index_map(sheet)
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if not rows:
        return error_response(400, "Excel 中没有可导入的数据")

    updated = 0
    unchanged = 0
    skipped = 0
    errors = []

    def base_team_query():
        query = db.query(Team)
        if module_key:
            query = query.filter(Team.module_key == module_key)
        return query

    for row_number, row in enumerate(rows, start=2):
        if not any(_normalize_excel_text(cell) for cell in row):
            continue

        team = None
        team_id_text = _normalize_excel_text(_first_cell_value(row, header_map, ["队伍ID", "队伍编号", "team_id", "ID"]))
        captain_sid = _normalize_excel_text(_first_cell_value(row, header_map, ["队长学号", "captain_student_id", "学号"]))
        team_name = _normalize_excel_text(_first_cell_value(row, header_map, ["队伍名称", "team_name"]))

        if team_id_text:
            try:
                team_id = int(float(team_id_text))
                team = base_team_query().filter(Team.id == team_id).first()
            except ValueError:
                errors.append(f"第 {row_number} 行：队伍ID格式不正确")

        if not team and captain_sid:
            team = base_team_query().filter(Team.captain_student_id == captain_sid).first()

        if not team and team_name:
            team = base_team_query().filter(Team.team_name == team_name).first()

        if not team:
            skipped += 1
            errors.append(f"第 {row_number} 行：未匹配到队伍")
            continue

        assignment = db.query(InspectionAssignment).filter(InspectionAssignment.team_id == team.id).first()
        if not assignment:
            assignment = InspectionAssignment(team_id=team.id, assigned_by=current_user.id)
            db.add(assignment)

        updates = {
            "first_inspection_time": _parse_excel_datetime(_first_cell_value(row, header_map, ["一验时间", "第一次验收时间", "first_inspection_time"])),
            "first_inspection_location": _normalize_excel_text(_first_cell_value(row, header_map, ["一验地点", "第一次验收地点", "first_inspection_location"])),
            "first_notes": _normalize_excel_text(_first_cell_value(row, header_map, ["一验备注", "第一次验收备注", "first_notes"])),
            "second_inspection_time": _parse_excel_datetime(_first_cell_value(row, header_map, ["二验时间", "第二次验收时间", "second_inspection_time"])),
            "second_inspection_location": _normalize_excel_text(_first_cell_value(row, header_map, ["二验地点", "第二次验收地点", "second_inspection_location"])),
            "second_notes": _normalize_excel_text(_first_cell_value(row, header_map, ["二验备注", "第二次验收备注", "second_notes"])),
        }

        row_changed = False
        for field_name, field_value in updates.items():
            if field_value in (None, ""):
                continue
            if getattr(assignment, field_name) != field_value:
                setattr(assignment, field_name, field_value)
                row_changed = True

        if row_changed:
            assignment.inspection_time = assignment.first_inspection_time
            assignment.inspection_location = assignment.first_inspection_location
            assignment.notes = assignment.first_notes
            assignment.assigned_by = current_user.id
            updated += 1
        else:
            unchanged += 1

    db.commit()
    return success_response(
        data={
            "updated": updated,
            "unchanged": unchanged,
            "skipped": skipped,
            "errors": errors[:20],
        },
        msg=f"导入完成，更新 {updated} 条，未变化 {unchanged} 条，跳过 {skipped} 条",
    )


@router.get("/teams/{team_id:int}", response_model=dict, summary="管理员查看队伍详情")
async def get_team_detail(
    team_id: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")
    return success_response(data=_team_to_dict(team), msg="获取成功")


@router.post("/teams/batch-delete", response_model=dict, summary="管理员批量删除队伍")
async def batch_delete_teams(
    req: TeamBatchDeleteRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team_ids = list(set(req.team_ids or []))
    if not team_ids:
        return error_response(400, "team_ids 不能为空")

    affected = db.query(Team).filter(Team.id.in_(team_ids), Team.delete_requested == True).delete(synchronize_session=False)
    db.commit()
    return success_response(data={"deleted": affected}, msg=f"确认删除成功，共 {affected} 条")


@router.post("/teams/batch-topic", response_model=dict, summary="管理员批量修改队伍选题")
async def batch_update_team_topic(
    req: TeamBatchTopicRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team_ids = list(set(req.team_ids or []))
    if not team_ids:
        return error_response(400, "team_ids 不能为空")

    topic_error = _assert_topic_available(db, req.topic_id)
    if topic_error:
        return error_response(400, topic_error)

    affected = db.query(Team).filter(Team.id.in_(team_ids)).update(
        {Team.topic_id: req.topic_id}, synchronize_session=False
    )
    db.commit()
    return success_response(data={"updated": affected}, msg=f"修改成功，共 {affected} 条")


@router.put("/teams/{team_id:int}/inspection", response_model=dict, summary="管理员更新验收安排")
async def update_inspection_assignment(
    team_id: int,
    req: TeamInspectionUpdateRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        return error_response(404, "队伍不存在")

    assignment = db.query(InspectionAssignment).filter(InspectionAssignment.team_id == team_id).first()
    if not assignment:
        assignment = InspectionAssignment(team_id=team_id, assigned_by=current_user.id)
        db.add(assignment)

    assignment.first_inspection_time = req.first_inspection_time
    assignment.first_inspection_location = req.first_inspection_location
    assignment.first_inspector = req.first_inspector
    assignment.first_notes = req.first_notes

    assignment.second_inspection_time = req.second_inspection_time
    assignment.second_inspection_location = req.second_inspection_location
    assignment.second_inspector = req.second_inspector
    assignment.second_notes = req.second_notes

    # 兼容旧字段：默认同步一验
    assignment.inspection_time = req.inspection_time if req.inspection_time is not None else req.first_inspection_time
    assignment.inspection_location = req.inspection_location if req.inspection_location is not None else req.first_inspection_location
    assignment.inspector = req.inspector if req.inspector is not None else req.first_inspector
    assignment.notes = req.notes if req.notes is not None else req.first_notes
    assignment.assigned_by = current_user.id

    db.commit()
    return success_response(msg="验收安排更新成功")


@router.post("/teams/batch-inspection", response_model=dict, summary="管理员批量更新验收安排")
async def batch_update_inspection_assignment(
    req: TeamBatchInspectionUpdateRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    team_ids = sorted(set(req.team_ids or []))
    if not team_ids:
        return error_response(400, "team_ids 不能为空")

    updatable_fields = [
        "first_inspection_time",
        "first_inspection_location",
        "first_inspector",
        "first_notes",
        "second_inspection_time",
        "second_inspection_location",
        "second_inspector",
        "second_notes",
        "inspection_time",
        "inspection_location",
        "inspector",
        "notes",
    ]

    payload = {}
    has_update = False
    for field in updatable_fields:
        value = getattr(req, field)
        if isinstance(value, str):
            value = value.strip()
            if value == "":
                value = None
        payload[field] = value
        if value is not None:
            has_update = True

    if not has_update:
        return error_response(400, "请至少提供一项验收信息用于更新")

    teams = db.query(Team).filter(Team.id.in_(team_ids)).all()
    found_ids = {team.id for team in teams}
    missing_ids = [team_id for team_id in team_ids if team_id not in found_ids]
    if missing_ids:
        return error_response(404, "存在不存在的队伍ID", data={"missing_team_ids": missing_ids})

    assignments = (
        db.query(InspectionAssignment)
        .filter(InspectionAssignment.team_id.in_(team_ids))
        .all()
    )
    assignment_map = {item.team_id: item for item in assignments}

    updated_count = 0
    skipped_duplicate_ids: List[int] = []

    for team in teams:
        assignment = assignment_map.get(team.id)
        if not assignment:
            assignment = InspectionAssignment(team_id=team.id, assigned_by=current_user.id)
            db.add(assignment)
            assignment_map[team.id] = assignment

        is_duplicate = True
        for field, value in payload.items():
            if value is None:
                continue
            current_value = getattr(assignment, field)
            if current_value != value:
                is_duplicate = False
                break

        if req.check_duplicate and is_duplicate:
            skipped_duplicate_ids.append(team.id)
            continue

        for field, value in payload.items():
            if value is not None:
                setattr(assignment, field, value)

        # 兼容旧字段：默认同步一验
        if req.inspection_time is None and req.first_inspection_time is not None:
            assignment.inspection_time = req.first_inspection_time
        if req.inspection_location is None and req.first_inspection_location is not None:
            assignment.inspection_location = req.first_inspection_location
        if req.inspector is None and req.first_inspector is not None:
            assignment.inspector = req.first_inspector
        if req.notes is None and req.first_notes is not None:
            assignment.notes = req.first_notes

        assignment.assigned_by = current_user.id
        updated_count += 1

    if updated_count == 0 and skipped_duplicate_ids:
        return error_response(
            400,
            "检测到重复安排，未执行更新",
            data={"skipped_duplicate_team_ids": skipped_duplicate_ids},
        )

    db.commit()

    msg = f"批量验收维护成功，共更新 {updated_count} 支队伍"
    if skipped_duplicate_ids:
        msg += f"，跳过重复安排 {len(skipped_duplicate_ids)} 支"

    return success_response(
        data={
            "updated": updated_count,
            "skipped_duplicate_team_ids": skipped_duplicate_ids,
        },
        msg=msg,
    )


@router.put("/teams/channel-config", response_model=dict, summary="管理员更新通道配置")
async def update_channel_config(
    req: TeamChannelConfigUpdateRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    cfg = _ensure_channel_config(db, req.module_key)

    if req.signup_open is not None:
        cfg.signup_open = req.signup_open
    if req.topic_open is not None:
        cfg.topic_open = req.topic_open
    if req.info_update_open is not None:
        cfg.info_update_open = req.info_update_open
    if req.signup_close_at is not None:
        cfg.signup_close_at = req.signup_close_at
    if req.topic_close_at is not None:
        cfg.topic_close_at = req.topic_close_at
    if req.info_update_close_at is not None:
        cfg.info_update_close_at = req.info_update_close_at
    cfg.updated_by = current_user.id

    db.commit()
    db.refresh(cfg)
    return success_response(
        data={
            "module_key": cfg.module_key,
            "signup_open": cfg.signup_open,
            "topic_open": cfg.topic_open,
            "info_update_open": cfg.info_update_open,
            "signup_close_at": cfg.signup_close_at.isoformat() if cfg.signup_close_at else None,
            "topic_close_at": cfg.topic_close_at.isoformat() if cfg.topic_close_at else None,
            "info_update_close_at": cfg.info_update_close_at.isoformat() if cfg.info_update_close_at else None,
        },
        msg="配置更新成功",
    )
